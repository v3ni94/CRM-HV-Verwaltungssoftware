"""P10: evaluations with common header, month matrix, target/actual, bank statement, Excel,
draft VAT overview, tax flags, rule version register, procedure documentation and the extended
audit export (7.7, 7.12; M18-01 to M18-09, SA-07, S711-11)."""

import asyncio
import csv
import hashlib
import io
import zipfile
from collections.abc import Iterator
from decimal import Decimal
from typing import Any

import boto3
import pytest
from fastapi.testclient import TestClient
from moto import mock_aws
from openpyxl import load_workbook

from mhvp.main import create_app
from mhvp.platform import services
from tests.integration.conftest import Database
from tests.integration.test_m2_platform import PASSWORD, RUN, World, _settings, bearer, login
from tests.integration.test_m18_audit_export import BUCKET
from tests.integration.test_m18_audit_export import _settings as _s3_settings
from tests.integration.test_m18_reports import _ok
from tests.integration.test_m18_tax_advisor_scope import assign_ledger_scope

pytestmark = pytest.mark.integration
A = "/api/v1/accounting"


async def _world(settings: Any) -> World:
    from mhvp.core import crypto
    from mhvp.core.db.engine import create_app_engine, create_session_factory

    crypto.set_master_key(b"k" * 32)
    engine = create_app_engine(settings)
    factory = create_session_factory(engine)
    try:
        a, _ = await services.provision_tenant(factory, slug=f"p10a-{RUN}", name=f"P10 A {RUN}")
        b, _ = await services.provision_tenant(factory, slug=f"p10b-{RUN}", name=f"P10 B {RUN}")
        world = World(tenant_a=a, tenant_b=b, app_url=settings.database_url.get_secret_value())
        for name, role, tenant in [
            ("p10admin", "tenant_admin", a),
            ("p10acc", "accountant_no_banking", a),
            ("p10tax", "tax_advisor", a),
            ("p10sup", "support", a),
            ("p10other", "tenant_admin", b),
        ]:
            uid = await services.create_user(
                factory, email=world.email(name), display_name=name, password=PASSWORD
            )
            world.users[name] = uid
            await services.add_member(
                factory, tenant_id=tenant, user_id=uid, role_codes=[role], actor_user_id=None
            )
        return world
    finally:
        await engine.dispose()


@pytest.fixture(scope="module")
def world(database: Database, redis_url: str) -> World:
    return asyncio.run(_world(_settings(database, redis_url)))


@pytest.fixture
def client(database: Database, redis_url: str) -> Iterator[TestClient]:
    with TestClient(create_app(_settings(database, redis_url))) as test_client:
        yield test_client


@pytest.fixture
def s3_client(database: Database, redis_url: str) -> Iterator[TestClient]:
    with mock_aws():
        boto3.client("s3", region_name="us-east-1").create_bucket(Bucket=BUCKET)
        with TestClient(create_app(_s3_settings(database, redis_url))) as test_client:
            yield test_client


def _setup(client: TestClient, h: dict[str, str], acc_user: dict[str, str]) -> dict[str, Any]:
    prop = _ok(
        client.post(
            "/api/v1/properties",
            json={"number": "791", "name": "P10 Haus", "management_type": "hoa"},
            headers=h,
        ),
        201,
    )
    hoa = next(e["id"] for e in prop["legal_entities"] if e["kind"] == "hoa")
    building = _ok(
        client.post(f"/api/v1/properties/{prop['id']}/buildings", json={"name": "Haus"}, headers=h),
        201,
    )["id"]
    unit = _ok(
        client.post(
            f"/api/v1/properties/{prop['id']}/units",
            json={"building_id": building, "number": "01", "unit_type": "apartment"},
            headers=h,
        ),
        201,
    )["id"]
    contact = _ok(
        client.post(
            "/api/v1/contacts",
            json={"kind": "person", "first_name": "Paul", "last_name": f"Z{RUN}"},
            headers=h,
        ),
        201,
    )
    party = _ok(
        client.post(
            "/api/v1/parties", json={"members": [{"contact_id": contact["id"]}]}, headers=h
        ),
        201,
    )["id"]
    contract = _ok(
        client.post(
            "/api/v1/contracts",
            json={
                "kind": "ownership",
                "unit_id": unit,
                "party_id": party,
                "start_date": "2020-01-01",
                "title_transfer_date": "2020-01-01",
                "acquisition_kind": "first_acquisition",
            },
            headers=h,
        ),
        201,
    )
    template = _ok(client.post(f"{A}/templates/default", headers=h), 201)
    ledger = _ok(
        client.post(
            f"{A}/ledgers", json={"legal_entity_id": hoa, "template_id": template["id"]}, headers=h
        ),
        201,
    )["id"]
    acc = {a["number"]: a for a in _ok(client.get(f"{A}/ledgers/{ledger}/accounts", headers=h))}
    debtor = next(a["id"] for a in acc.values() if a["category"] == "debtor")

    def book(body: dict[str, Any]) -> dict[str, Any]:
        draft = _ok(client.post(f"{A}/ledgers/{ledger}/entries", json=body, headers=h), 201)
        if body["kind"] == "opening_balance":
            _ok(
                client.post(f"{A}/ledgers/{ledger}/entries/{draft['id']}/approve", headers=acc_user)
            )
        return _ok(client.post(f"{A}/ledgers/{ledger}/entries/{draft['id']}/post", headers=h))  # type: ignore[no-any-return]

    book(
        {
            "kind": "opening_balance",
            "booking_date": "2026-01-01",
            "text": "Anfangsbestand",
            "lines": [
                {"account_id": acc["001200"]["id"], "debit": "5000.00"},
                {"account_id": acc["001201"]["id"], "debit": "20000.00"},
                {"account_id": acc["009000"]["id"], "credit": "25000.00"},
            ],
        }
    )
    book(
        {
            "kind": "receivable",
            "booking_date": "2026-09-01",
            "due_date": "2026-09-03",
            "text": "Hausgeld",
            "lines": [
                {"account_id": debtor, "debit": "400.00"},
                {"account_id": acc["060100"]["id"], "credit": "400.00"},
            ],
        }
    )
    oi = _ok(
        client.get(f"{A}/ledgers/{ledger}/open-items", params={"as_of": "2026-09-30"}, headers=h)
    )[0]["id"]
    book(
        {
            "kind": "debtor_payment",
            "booking_date": "2026-09-05",
            "text": "Zahlung",
            "settlements": [{"open_item_id": oi, "amount": "150.00"}],
            "lines": [
                {"account_id": acc["001200"]["id"], "debit": "150.00"},
                {"account_id": debtor, "credit": "150.00"},
            ],
        }
    )
    return {"ledger": ledger, "acc": acc, "debtor": debtor, "contract": contract}


@pytest.fixture(scope="module")
def booked(database: Database, redis_url: str, world: World) -> dict[str, Any]:
    with TestClient(create_app(_settings(database, redis_url))) as c:
        h = bearer(login(c, world, "p10admin"))
        acc_user = bearer(login(c, world, "p10acc"))
        data = _setup(c, h, acc_user)
        assign_ledger_scope(c, h, world.users["p10tax"], data["ledger"])
        return data


P = {"start": "2026-01-01", "end": "2026-12-31"}


def test_reports_with_header(client: TestClient, world: World, booked: dict[str, Any]) -> None:
    h = bearer(login(client, world, "p10admin"))
    ledger = booked["ledger"]
    matrix = _ok(client.get(f"{A}/ledgers/{ledger}/reports/monthly-matrix", params=P, headers=h))
    head = matrix["header"]
    assert head["status"] == "draft"
    assert head["legal_entity_name"]
    assert head["period_start"] == "2026-01-01"
    assert head["generated_at"]
    assert len(matrix["months"]) == 12
    row = next(a for a in matrix["accounts"] if a["number"] == "060100")
    assert Decimal(row["months"]["2026-09"]) == Decimal("400.00")
    assert Decimal(row["total"]) == Decimal("400.00")
    assert (
        client.get(
            f"{A}/ledgers/{ledger}/reports/monthly-matrix",
            params={"start": "2026-12-31", "end": "2026-01-01"},
            headers=h,
        ).status_code
        == 422
    )

    ta = _ok(client.get(f"{A}/ledgers/{ledger}/reports/target-actual", params=P, headers=h))
    assert Decimal(ta["total_target"]) == Decimal("400.00")
    assert Decimal(ta["total_actual_on_target"]) == Decimal("150.00")
    assert Decimal(ta["total_difference"]) == Decimal("250.00")
    assert Decimal(ta["total_receipts_in_period"]) == Decimal("150.00")
    assert ta["header"]["report"] == "target_actual"

    bank = _ok(
        client.get(
            f"{A}/ledgers/{ledger}/reports/bank-statement",
            params={
                "account_id": booked["acc"]["001200"]["id"],
                "start": "2026-02-01",
                "end": "2026-12-31",
            },
            headers=h,
        )
    )
    assert Decimal(bank["opening_balance"]) == Decimal("5000.00")
    assert Decimal(bank["closing_balance"]) == Decimal("5150.00")
    assert len(bank["movements"]) == 1
    assert bank["reconciliation"] is None
    assert (
        client.get(
            f"{A}/ledgers/{ledger}/reports/bank-statement",
            params={"account_id": booked["debtor"], **P},
            headers=h,
        ).status_code
        == 404
    )

    tb = _ok(
        client.get(
            f"{A}/ledgers/{ledger}/reports/trial-balance", params={"as_of": "2026-09-30"}, headers=h
        )
    )
    assert tb["balanced"] is True
    assert tb["header"]["as_of"] == "2026-09-30"
    ops = _ok(
        client.get(
            f"{A}/ledgers/{ledger}/reports/open-items", params={"as_of": "2026-09-30"}, headers=h
        )
    )
    assert [Decimal(r["remaining"]) for r in ops["rows"]] == [Decimal("250.00")]
    sheet = _ok(
        client.get(
            f"{A}/ledgers/{ledger}/reports/account-sheet",
            params={"account_id": booked["acc"]["001200"]["id"], **P},
            headers=h,
        )
    )
    assert sheet["header"]["filters"]["account"].startswith("001200")
    rev = _ok(client.get(f"{A}/ledgers/{ledger}/reports/revenue", params=P, headers=h))
    assert [r["number"] for r in rev["rows"]] == ["060100"]
    ie = _ok(client.get(f"{A}/ledgers/{ledger}/reports/income-expense", params=P, headers=h))
    assert Decimal(ie["total_revenue"]) == Decimal("400.00")
    assert "keine Einnahmenüberschussrechnung" in ie["note"]


def test_vat_overview_and_tax_flags(
    client: TestClient, world: World, booked: dict[str, Any]
) -> None:
    h = bearer(login(client, world, "p10admin"))
    tax = bearer(login(client, world, "p10tax"))
    ledger, acc = booked["ledger"], booked["acc"]
    vat = _ok(client.get(f"{A}/ledgers/{ledger}/reports/vat-overview", params=P, headers=tax))
    assert vat["header"]["status"] == "draft"
    assert Decimal(vat["total_output_vat"]) == 0
    assert {c["id"] for c in vat["checkpoints"]} >= {"S711-03", "S711-05"}
    assert all(c["status"] == "offen" for c in vat["checkpoints"])
    assert "keine Umsatzsteuer-Voranmeldung" in vat["note"]

    url = f"{A}/ledgers/{ledger}/accounts/{acc['060100']['id']}/tax-flags"
    body = {"eur_relevant": True, "ust_relevant": True, "mixed_use_review": True}
    assert client.put(url, json=body, headers=tax).status_code == 403
    assert client.put(url, json={"eur_relevant": True}, headers=h).status_code == 422
    out = _ok(client.put(url, json=body, headers=h))
    assert out["eur_relevant"]
    assert out["ust_relevant"]
    assert out["mixed_use_review"]
    vat = _ok(client.get(f"{A}/ledgers/{ledger}/reports/vat-overview", params=P, headers=h))
    assert vat["ust_flagged_accounts"] == 1
    assert vat["mixed_use_review_accounts"] == [
        {"number": "060100", "name": vat["mixed_use_review_accounts"][0]["name"]}
    ]
    matrix = _ok(
        client.get(
            f"{A}/ledgers/{ledger}/reports/monthly-matrix",
            params={**P, "eur_only": "true"},
            headers=h,
        )
    )
    assert [a["number"] for a in matrix["accounts"]] == ["060100"]
    accounts = _ok(client.get(f"{A}/ledgers/{ledger}/accounts", headers=h))
    assert next(a for a in accounts if a["number"] == "060100")["ust_relevant"] is True


def test_xlsx_and_procedure_documentation(
    client: TestClient, world: World, booked: dict[str, Any]
) -> None:
    tax = bearer(login(client, world, "p10tax"))
    sup = bearer(login(client, world, "p10sup"))
    ledger = booked["ledger"]
    response = client.get(
        f"{A}/ledgers/{ledger}/reports/xlsx", params={"report": "journal", **P}, headers=tax
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["x-content-sha256"] == hashlib.sha256(response.content).hexdigest()
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    values = [[c.value for c in row] for row in sheet.iter_rows()]
    assert values[0][:2] == ["Auswertung", "journal"]
    assert ["Status", "Entwurf"] in [row[:2] for row in values]
    header_index = next(i for i, row in enumerate(values) if row[0] == "Jahr")
    assert len(values) - header_index - 1 == 7
    for report in (
        "monthly_matrix",
        "target_actual",
        "trial_balance",
        "open_items",
        "vat_overview",
    ):
        r = client.get(
            f"{A}/ledgers/{ledger}/reports/xlsx",
            params={"report": report, "as_of": "2026-09-30", **P},
            headers=tax,
        )
        assert r.status_code == 200, (report, r.text)
    assert (
        client.get(
            f"{A}/ledgers/{ledger}/reports/xlsx", params={"report": "journal", **P}, headers=sup
        ).status_code
        == 403
    )
    assert (
        client.get(
            f"{A}/ledgers/{ledger}/reports/xlsx", params={"report": "unknown", **P}, headers=tax
        ).status_code
        == 422
    )

    doc = client.get(f"{A}/ledgers/{ledger}/procedure-documentation", headers=tax)
    assert doc.status_code == 200, doc.text
    assert doc.headers["x-draft"] == "true"
    assert "Entwurf" in doc.text
    assert "Gebuchte Sätze: 3" in doc.text
    assert "G1 Produktive Buchführung: gesperrt" in doc.text
    download = client.get(
        f"{A}/ledgers/{ledger}/procedure-documentation", params={"download": "true"}, headers=tax
    )
    assert "attachment" in download.headers["content-disposition"]
    assert (
        client.get(f"{A}/ledgers/{ledger}/procedure-documentation", headers=sup).status_code == 403
    )


def test_tenant_separation_and_permissions(
    client: TestClient, world: World, booked: dict[str, Any]
) -> None:
    other = bearer(login(client, world, "p10other"))
    sup = bearer(login(client, world, "p10sup"))
    ledger = booked["ledger"]
    for path, params in (
        ("reports/monthly-matrix", P),
        ("reports/target-actual", P),
        ("reports/vat-overview", P),
        ("reports/trial-balance", {"as_of": "2026-09-30"}),
        ("procedure-documentation", {}),
    ):
        assert (
            client.get(f"{A}/ledgers/{ledger}/{path}", params=params, headers=other).status_code
            == 404
        ), path
        assert (
            client.get(f"{A}/ledgers/{ledger}/{path}", params=params, headers=sup).status_code
            == 403
        ), path


def test_rule_version_register(client: TestClient, world: World) -> None:
    h = bearer(login(client, world, "p10admin"))
    tax = bearer(login(client, world, "p10tax"))
    other = bearer(login(client, world, "p10other"))
    body = {
        "rule_id": "S711-TEST",
        "title": "Testregel",
        "effective_from": "2026-01-01",
        "case_groups": ["weg", "miete"],
        "source_status": "zu prüfen",
        "change_reason": "Neuregelung",
    }
    assert client.post(f"{A}/rule-versions", json=body, headers=tax).status_code == 403
    assert (
        client.post(f"{A}/rule-versions", json={**body, "rule_id": "bad id"}, headers=h).status_code
        == 422
    )
    assert (
        client.post(
            f"{A}/rule-versions", json={**body, "effective_to": "2025-01-01"}, headers=h
        ).status_code
        == 422
    )
    v1 = _ok(client.post(f"{A}/rule-versions", json=body, headers=h), 201)
    assert v1["version"] == 1
    assert v1["status"] == "draft"
    early = client.post(
        f"{A}/rule-versions", json={**body, "effective_from": "2025-06-01"}, headers=h
    )
    assert early.status_code == 422
    v2 = _ok(
        client.post(f"{A}/rule-versions", json={**body, "effective_from": "2027-01-01"}, headers=h),
        201,
    )
    assert v2["version"] == 2
    eff = _ok(
        client.get(
            f"{A}/rule-versions/effective",
            params={"rule_id": "S711-TEST", "on": "2026-06-30"},
            headers=tax,
        )
    )
    assert eff["version"] == 1
    eff = _ok(
        client.get(
            f"{A}/rule-versions/effective",
            params={"rule_id": "S711-TEST", "on": "2027-02-01"},
            headers=tax,
        )
    )
    assert eff["version"] == 2
    assert (
        client.get(
            f"{A}/rule-versions/effective",
            params={"rule_id": "S711-TEST", "on": "2025-01-01"},
            headers=tax,
        ).status_code
        == 404
    )
    confirmed = _ok(
        client.post(
            f"{A}/rule-versions/{v1['id']}/confirm",
            json={"confirmed_by": "StB Beispiel", "confirmed_on": "2026-09-30"},
            headers=h,
        )
    )
    assert confirmed["status"] == "confirmed"
    assert confirmed["expert_confirmed_by"] == "StB Beispiel"
    again = client.post(
        f"{A}/rule-versions/{v1['id']}/confirm",
        json={"confirmed_by": "StB Beispiel", "confirmed_on": "2026-09-30"},
        headers=h,
    )
    assert again.status_code == 409
    withdrawn = _ok(client.post(f"{A}/rule-versions/{v2['id']}/withdraw", headers=h))
    assert withdrawn["status"] == "withdrawn"
    listed = _ok(client.get(f"{A}/rule-versions", params={"rule_id": "S711-TEST"}, headers=tax))
    assert [r["version"] for r in listed] == [1, 2]
    # Other tenant sees nothing and cannot touch the rows.
    assert (
        _ok(client.get(f"{A}/rule-versions", params={"rule_id": "S711-TEST"}, headers=other)) == []
    )
    assert client.post(f"{A}/rule-versions/{v1['id']}/withdraw", headers=other).status_code == 404


def test_audit_export_contracts_and_approvals(
    s3_client: TestClient, world: World, booked: dict[str, Any]
) -> None:
    tax = bearer(login(s3_client, world, "p10tax"))
    run = _ok(
        s3_client.post(
            f"{A}/audit-exports",
            json={
                "ledger_id": booked["ledger"],
                "period_from": "2026-01-01",
                "period_to": "2026-12-31",
            },
            headers=tax,
        ),
        201,
    )
    data = s3_client.get(f"{A}/audit-exports/{run['id']}/download", headers=tax).content
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert "vertraege.csv" in zf.namelist()
        text = zf.read("vertraege.csv").decode("utf-8").removeprefix("﻿")
        rows = list(csv.DictReader(io.StringIO(text), delimiter=";"))
        assert [r["Nummer"] for r in rows] == [booked["contract"]["number"]]
        assert rows[0]["Version"] == "1"
        assert rows[0]["Art"] == "ownership"
        assert rows[0]["Beginn"] == "2020-01-01"
