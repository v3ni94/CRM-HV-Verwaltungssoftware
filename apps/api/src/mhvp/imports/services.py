"""Import assistant (13.1): read export files, stage rows, map and validate, test run, apply
as ``import_run`` and reconcile.

Existing records are never overwritten: a row that matches an existing record with the same
values is ``unchanged``, with different values ``conflict`` (for manual review).
"""

import csv
import io
import uuid
from collections import Counter
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from mhvp.ai.imports import Recorder, create_contact, create_party
from mhvp.ai.models import ImportRun
from mhvp.contacts import schemas as cs
from mhvp.contacts.models import Contact, Party, PartyMember
from mhvp.contracts import schemas as contract_schemas
from mhvp.contracts import services as contract_services
from mhvp.contracts.models import Contract, ContractKind, ContractPayment, PaymentReason
from mhvp.core.problems import ErrorCodes, ProblemError
from mhvp.imports.fields import FIELDS, convert
from mhvp.imports.models import ImportSourceFile, ReportType, RowStatus, StagingRow
from mhvp.properties import services as property_services
from mhvp.properties.models import (
    AllocationKey,
    Building,
    ManagementType,
    Property,
    PropertyOwner,
    PropertyStatus,
    Unit,
    UnitType,
    ValueSource,
)

MAX_ROWS = 20_000
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CENT = Decimal("0.01")


def _cell(value: Any) -> Any:
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, Decimal):
        return str(value)
    return value


def read_table(
    data: bytes, mime_type: str, sheet: str | None, header_row: int
) -> tuple[list[str], list[dict[str, Any]]]:
    if mime_type == XLSX:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        if sheet is not None and sheet not in book.sheetnames:
            raise ProblemError(ErrorCodes.VALIDATION, detail=f"Tabellenblatt {sheet!r} fehlt.")
        rows = [list(r) for r in book[sheet or book.sheetnames[0]].iter_rows(values_only=True)]
    elif mime_type in ("text/csv", "text/plain"):
        text = data.decode("utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ProblemError(ErrorCodes.VALIDATION, detail="Nur Excel (xlsx) oder CSV.")
    if len(rows) < header_row:
        raise ProblemError(ErrorCodes.VALIDATION, detail="Die Kopfzeile fehlt.")
    headers = [str(h).strip() if h is not None else "" for h in rows[header_row - 1]]
    duplicates = [h for h, n in Counter(h for h in headers if h).items() if n > 1]
    if duplicates:
        raise ProblemError(
            ErrorCodes.VALIDATION, detail=f"Doppelte Spaltenüberschriften: {', '.join(duplicates)}."
        )
    body = []
    for row in rows[header_row:]:
        if all(c in (None, "") for c in row):
            continue
        body.append(
            {headers[i]: _cell(c) for i, c in enumerate(row) if i < len(headers) and headers[i]}
        )
    if len(body) > MAX_ROWS:
        raise ProblemError(
            ErrorCodes.VALIDATION, detail=f"Mehr als {MAX_ROWS} Zeilen; bitte teilen."
        )
    return [h for h in headers if h], body


# Lookups ---------------------------------------------------------------------------------


async def _property(session: AsyncSession, number: str) -> Property | None:
    row: Property | None = await session.scalar(select(Property).where(Property.number == number))
    return row


async def _unit(session: AsyncSession, property_number: str, number: str) -> Unit | None:
    row: Unit | None = await session.scalar(
        select(Unit)
        .join(Property, Property.id == Unit.property_id)
        .where(Property.number == property_number, Unit.number == number)
    )
    return row


async def _contact(session: AsyncSession, external_id: str) -> Contact | None:
    row: Contact | None = await session.scalar(
        select(Contact).where(
            Contact.external_ids["immoware24"].astext == external_id, Contact.deleted_at.is_(None)
        )
    )
    return row


async def _party_of(session: AsyncSession, contact: Contact) -> Party | None:
    """The party whose only member is this contact (created by the contacts import)."""
    single = select(PartyMember.party_id).group_by(PartyMember.party_id).having(func.count() == 1)
    row: Party | None = await session.scalar(
        select(Party)
        .join(PartyMember, PartyMember.party_id == Party.id)
        .where(PartyMember.contact_id == contact.id, Party.id.in_(single))
        .order_by(Party.created_at)
        .limit(1)
    )
    return row


# Row handlers: return (status, entity_type, entity_id, errors) ---------------------------


Result = tuple[RowStatus, str | None, uuid.UUID | None, list[str]]


def _same(existing: Any, values: dict[str, Any], fields: tuple[str, ...]) -> bool:
    return all(
        values.get(f) is None or str(getattr(existing, f) or "") == str(values[f]) for f in fields
    )


async def _apply_property(
    session: AsyncSession, principal: Any, v: dict[str, Any], rec: Recorder | None
) -> Result:
    existing = await _property(session, v["number"])
    if existing is not None:
        keys = ("name", "street", "house_number", "postal_code", "city")
        same = _same(existing, v, keys) and existing.management_type.value == v["management_type"]
        return (
            (RowStatus.UNCHANGED if same else RowStatus.CONFLICT),
            "property",
            existing.id,
            ([] if same else ["Objekt vorhanden mit abweichenden Angaben"]),
        )
    prop = Property(
        tenant_id=principal.tenant_id,
        created_by=principal.user_id,
        number=v["number"],
        name=v["name"],
        management_type=ManagementType(v["management_type"]),
        street=v.get("street"),
        house_number=v.get("house_number"),
        postal_code=v.get("postal_code"),
        city=v.get("city"),
        status=PropertyStatus.ONBOARDING,
    )
    session.add(prop)
    await session.flush()
    await property_services.ensure_hoa_entity(session, prop)
    await property_services.copy_key_templates(session, prop)
    if rec:
        rec.add("property", prop.id)
    return RowStatus.CREATED, "property", prop.id, []


async def _apply_unit(
    session: AsyncSession, principal: Any, v: dict[str, Any], rec: Recorder | None
) -> Result:
    prop = await _property(session, v["property_number"])
    if prop is None:
        return RowStatus.INVALID, None, None, [f"Objekt {v['property_number']} fehlt"]
    if v.get("mea") and not v.get("mea_valid_from"):
        return RowStatus.INVALID, None, None, ["MEA ohne Gültig-ab-Datum"]
    existing = await _unit(session, v["property_number"], v["number"])
    if existing is not None:
        same = (
            _same(existing, v, ("label", "location"))
            and existing.unit_type.value == v["unit_type"]
            and (
                v.get("living_area_sqm") is None
                or Decimal(v["living_area_sqm"]) == (existing.living_area_sqm or Decimal(-1))
            )
        )
        return (
            (RowStatus.UNCHANGED if same else RowStatus.CONFLICT),
            "unit",
            existing.id,
            ([] if same else ["Einheit vorhanden mit abweichenden Angaben"]),
        )
    name = v.get("building") or prop.name
    building = await session.scalar(
        select(Building).where(Building.property_id == prop.id, Building.name == name)
    )
    if building is None:
        building = Building(tenant_id=principal.tenant_id, property_id=prop.id, name=name[:200])
        session.add(building)
        await session.flush()
        if rec:
            rec.add("building", building.id)
    unit = Unit(
        tenant_id=principal.tenant_id,
        property_id=prop.id,
        building_id=building.id,
        number=v["number"][:20],
        label=v.get("label"),
        location=v.get("location"),
        unit_type=UnitType(v["unit_type"]),
        living_area_sqm=Decimal(v["living_area_sqm"]) if v.get("living_area_sqm") else None,
    )
    session.add(unit)
    await session.flush()
    if rec:
        rec.add("unit", unit.id)
    if v.get("mea"):
        key = await session.scalar(
            select(AllocationKey).where(
                AllocationKey.property_id == prop.id, AllocationKey.code == "MEA"
            )
        )
        if key is not None:
            await property_services.add_allocation_value(
                session,
                principal.tenant_id,
                unit.id,
                key.id,
                Decimal(v["mea"]),
                date.fromisoformat(v["mea_valid_from"]),
                None,
                ValueSource.IMPORT,
            )
    return RowStatus.CREATED, "unit", unit.id, []


async def _apply_contact(
    session: AsyncSession, principal: Any, v: dict[str, Any], rec: Recorder | None
) -> Result:
    existing = await _contact(session, v["external_id"])
    if existing is not None:
        same = _same(existing, v, ("first_name", "last_name", "company_name"))
        return (
            (RowStatus.UNCHANGED if same else RowStatus.CONFLICT),
            "contact",
            existing.id,
            ([] if same else ["Kontakt vorhanden mit abweichendem Namen"]),
        )
    notes: list[str] = []
    data: dict[str, Any] = {
        "kind": v["kind"],
        "salutation": v.get("salutation"),
        "title": v.get("title"),
        "first_name": v.get("first_name"),
        "last_name": v.get("last_name"),
        "company_name": v.get("company_name"),
        "external_ids": {"immoware24": v["external_id"]},
    }
    address = {k: v.get(k) for k in ("street", "house_number", "postal_code", "city")}
    if any(address.values()):
        data["addresses"] = [{**address, "is_primary": True}]
    for key, schema, target in (("phone", cs.PhoneIn, "phones"), ("email", cs.EmailIn, "emails")):
        if v.get(key):
            try:
                field = "number" if key == "phone" else "email"
                data[target] = [schema.model_validate({field: v[key]}).model_dump()]
            except ValueError:
                notes.append(f"{key} ungültig, nicht übernommen")
    if v.get("iban"):
        try:
            # The export carries no validity date: the account counts from the import day (A-021).
            today = datetime.now(UTC).date()
            data["bank_accounts"] = [
                cs.BankAccountIn(iban=v["iban"], valid_from=today).model_dump()
            ]
        except ValueError:
            notes.append("IBAN ungültig, nicht übernommen")
    try:
        contact_in = cs.ContactIn.model_validate(data)
    except ValueError as exc:
        return RowStatus.INVALID, None, None, [f"Kontakt ungültig: {exc}"]
    contact = await create_contact(session, principal.tenant_id, principal.user_id, contact_in)
    party = await create_party(session, principal.tenant_id, principal.user_id, [contact])
    if rec:
        rec.add("contact", contact.id)
        rec.add("party", party.id)
    return RowStatus.CREATED, "contact", contact.id, notes


async def _apply_contract(
    session: AsyncSession,
    principal: Any,
    v: dict[str, Any],
    rec: Recorder | None,
    kind: ContractKind,
) -> Result:
    from mhvp.contracts.routers import _create

    unit = await _unit(session, v["property_number"], v["unit_number"])
    contact = await _contact(session, v["contact_external_id"])
    if unit is None or contact is None:
        return (
            RowStatus.INVALID,
            None,
            None,
            [m for m, ok in (("Einheit fehlt", unit), ("Kontakt fehlt", contact)) if not ok],
        )
    party = await _party_of(session, contact)
    if party is None:
        return RowStatus.INVALID, None, None, ["Keine Vertragspartei zum Kontakt"]
    start = date.fromisoformat(v["start_date"])
    prop = await session.get(Property, unit.property_id)
    assert prop is not None  # noqa: S101 - unit references it
    if kind is ContractKind.OWNERSHIP and prop.management_type is ManagementType.RENTAL:
        existing_owner = await session.scalar(
            select(PropertyOwner).where(
                PropertyOwner.property_id == prop.id, PropertyOwner.party_id == party.id
            )
        )
        if existing_owner is not None:
            return RowStatus.UNCHANGED, "property_owner", existing_owner.id, []
        owner = PropertyOwner(
            tenant_id=principal.tenant_id, property_id=prop.id, party_id=party.id, valid_from=start
        )
        session.add(owner)
        await session.flush()
        await property_services.owner_entity(session, prop, party.id)
        if rec:
            rec.add("property_owner", owner.id)
        return RowStatus.CREATED, "property_owner", owner.id, []
    existing = await session.scalar(
        select(Contract).where(
            Contract.unit_id == unit.id, Contract.kind == kind, Contract.start_date == start
        )
    )
    if existing is not None:
        same = existing.party_id == party.id
        return (
            (RowStatus.UNCHANGED if same else RowStatus.CONFLICT),
            "contract",
            existing.id,
            ([] if same else ["Vertrag mit anderer Partei vorhanden"]),
        )
    body = contract_schemas.ContractIn(
        kind=kind,
        unit_id=unit.id,
        party_id=party.id,
        start_date=start,
        end_date=date.fromisoformat(v["end_date"]) if v.get("end_date") else None,
        title_transfer_date=date.fromisoformat(v["title_transfer_date"])
        if v.get("title_transfer_date")
        else None,
        benefit_burden_date=date.fromisoformat(v["benefit_burden_date"])
        if v.get("benefit_burden_date")
        else None,
    )
    try:
        async with session.begin_nested():
            contract = await _create(session, principal, body)
    except ProblemError as exc:
        return RowStatus.INVALID, None, None, [str(exc.detail)]
    if rec:
        rec.add("contract", contract.id)
    return RowStatus.CREATED, "contract", contract.id, []


async def _apply_payment(
    session: AsyncSession, principal: Any, v: dict[str, Any], rec: Recorder | None
) -> Result:
    unit = await _unit(session, v["property_number"], v["unit_number"])
    contact = await _contact(session, v["contact_external_id"])
    party = await _party_of(session, contact) if contact else None
    if unit is None or party is None:
        return RowStatus.INVALID, None, None, ["Einheit oder Vertragspartei fehlt"]
    valid_from = date.fromisoformat(v["valid_from"])
    contract = await session.scalar(
        select(Contract)
        .where(
            Contract.unit_id == unit.id,
            Contract.kind == ContractKind(v["kind"]),
            Contract.party_id == party.id,
            Contract.start_date <= valid_from,
        )
        .order_by(Contract.start_date.desc())
        .limit(1)
    )
    if contract is None:
        return RowStatus.INVALID, None, None, ["Kein passender Vertrag zum Gültig-ab-Datum"]
    code, gross, rate = v["payment_type_code"], Decimal(v["gross"]), Decimal(v["vat_percent"])
    existing = await session.scalar(
        select(ContractPayment).where(
            ContractPayment.contract_id == contract.id,
            ContractPayment.payment_type_code == code,
            ContractPayment.valid_from == valid_from,
        )
    )
    if existing is not None:
        same = existing.gross == gross
        return (
            (RowStatus.UNCHANGED if same else RowStatus.CONFLICT),
            "contract_payment",
            existing.id,
            ([] if same else ["Zahlung mit anderem Betrag vorhanden"]),
        )
    net = (gross / (1 + rate / 100)).quantize(CENT)
    try:
        contract_services.check_amounts(code, net, rate, gross)
        row = ContractPayment(
            tenant_id=principal.tenant_id,
            contract_id=contract.id,
            payment_type_code=code,
            net=net,
            vat_percent=rate,
            gross=gross,
            valid_from=valid_from,
            reason=PaymentReason.INITIAL,
        )
        async with session.begin_nested():
            await contract_services.add_payment(session, contract, row)
            await session.flush()
    except ProblemError as exc:
        return RowStatus.INVALID, None, None, [str(exc.detail)]
    if rec:
        rec.add("contract_payment", row.id)
    return RowStatus.CREATED, "contract_payment", row.id, []


async def apply_row(
    session: AsyncSession,
    principal: Any,
    report_type: ReportType,
    values: dict[str, Any],
    rec: Recorder | None,
) -> Result:
    if report_type is ReportType.PROPERTIES:
        return await _apply_property(session, principal, values, rec)
    if report_type is ReportType.UNITS:
        return await _apply_unit(session, principal, values, rec)
    if report_type is ReportType.CONTACTS:
        return await _apply_contact(session, principal, values, rec)
    if report_type is ReportType.TENANCIES:
        return await _apply_contract(session, principal, values, rec, ContractKind.TENANCY)
    if report_type is ReportType.OWNERSHIPS:
        return await _apply_contract(session, principal, values, rec, ContractKind.OWNERSHIP)
    if report_type is ReportType.PAYMENTS:
        return await _apply_payment(session, principal, values, rec)
    return RowStatus.STAGED_ONLY, None, None, []


# Runs ------------------------------------------------------------------------------------


async def validate(
    session: AsyncSession,
    source: ImportSourceFile,
    columns: dict[str, str],
    value_maps: dict[str, dict[str, str]],
) -> dict[str, int]:
    unknown = sorted(set(columns.values()) - set(source.headers))
    if unknown:
        raise ProblemError(
            ErrorCodes.VALIDATION, detail=f"Spalten fehlen in der Datei: {', '.join(unknown)}."
        )
    missing = [f.name for f in FIELDS[source.report_type] if f.required and f.name not in columns]
    if missing:
        raise ProblemError(
            ErrorCodes.VALIDATION, detail=f"Pflichtfelder ohne Zuordnung: {', '.join(missing)}."
        )
    rows = (
        await session.scalars(select(StagingRow).where(StagingRow.source_file_id == source.id))
    ).all()
    counts: Counter[str] = Counter()
    for row in rows:
        if not FIELDS[source.report_type]:
            row.values, row.errors, row.status = row.raw, [], RowStatus.STAGED_ONLY
        else:
            values, errors = convert(source.report_type, row.raw, columns, value_maps)
            row.values, row.errors = values, errors
            row.status = RowStatus.INVALID if errors else RowStatus.VALID
        counts[row.status.value] += 1
    await session.flush()
    return dict(counts)


async def run(
    session: AsyncSession, principal: Any, source: ImportSourceFile, import_run: ImportRun | None
) -> dict[str, Any]:
    """Apply valid rows; with ``import_run=None`` this is a test run the caller rolls back."""
    rec = Recorder(session, import_run) if import_run is not None else None
    rows = (
        await session.scalars(
            select(StagingRow)
            .where(StagingRow.source_file_id == source.id, StagingRow.status == RowStatus.VALID)
            .order_by(StagingRow.row_number)
        )
    ).all()
    counts: Counter[str] = Counter()
    errors: list[dict[str, Any]] = []
    source_sum = created_sum = Decimal(0)
    for row in rows:
        assert row.values is not None  # noqa: S101 - validated rows carry values
        status, entity_type, entity_id, messages = await apply_row(
            session, principal, source.report_type, row.values, rec
        )
        counts[status.value] += 1
        if source.report_type is ReportType.PAYMENTS:
            source_sum += Decimal(row.values["gross"])
            if status is RowStatus.CREATED:
                created_sum += Decimal(row.values["gross"])
        if messages:
            errors.append({"row": row.row_number, "status": status.value, "messages": messages})
        if import_run is not None:
            row.status, row.entity_type, row.entity_id = status, entity_type, entity_id
            row.errors = messages
    report: dict[str, Any] = {"counts": dict(counts), "problems": errors[:500]}
    if source.report_type is ReportType.PAYMENTS:
        report["sums"] = {"source_gross": str(source_sum), "created_gross": str(created_sum)}
    await session.flush()
    return report


async def reconcile(session: AsyncSession, source: ImportSourceFile) -> dict[str, Any]:
    """Compare the file with the platform after import (13.1 Abgleichbericht)."""
    rows = (
        await session.scalars(select(StagingRow).where(StagingRow.source_file_id == source.id))
    ).all()
    by_status = Counter(r.status.value for r in rows)
    result: dict[str, Any] = {"rows": len(rows), "status": dict(by_status)}
    if source.report_type is ReportType.UNITS:
        expected = Counter(
            r.values["property_number"] for r in rows if r.values and "property_number" in r.values
        )
        actual = {}
        for number in expected:
            actual[number] = (
                await session.scalar(
                    select(func.count())
                    .select_from(Unit)
                    .join(Property, Property.id == Unit.property_id)
                    .where(Property.number == number)
                )
                or 0
            )
        result["units_per_property"] = {
            n: {"file": expected[n], "platform": actual[n], "difference": actual[n] - expected[n]}
            for n in sorted(expected)
        }
    open_rows = by_status.get("invalid", 0) + by_status.get("conflict", 0)
    result["open_differences"] = open_rows
    return result
