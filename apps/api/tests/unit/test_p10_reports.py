"""P10: pure parts of the evaluations: month keys, Excel cell safety, header block and the
procedure documentation text."""

import io
from datetime import UTC, date, datetime
from decimal import Decimal
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from mhvp.accounting import procedure_doc, report_views, report_xlsx
from mhvp.accounting.models import AccountCategory, LeadingSystem, LedgerAccount, VatMode
from mhvp.core.problems import ProblemError


def test_month_keys_cover_the_period() -> None:
    assert report_views.month_keys(date(2026, 11, 15), date(2027, 2, 1)) == [
        "2026-11",
        "2026-12",
        "2027-01",
        "2027-02",
    ]


def test_month_keys_reject_reversed_and_long_periods() -> None:
    with pytest.raises(ProblemError):
        report_views.month_keys(date(2026, 2, 1), date(2026, 1, 1))
    with pytest.raises(ProblemError):
        report_views.month_keys(date(2020, 1, 1), date(2026, 1, 1))


def test_tax_flags_default_to_false() -> None:
    columns = LedgerAccount.__table__.c
    for name in ("eur_relevant", "ust_relevant", "mixed_use_review"):
        assert columns[name].nullable is False
        assert str(columns[name].server_default.arg) == "false"  # type: ignore[union-attr]


def test_checkpoints_are_open_without_assessment() -> None:
    ids = {c["id"]: c for c in report_views.TAX_CHECKPOINTS}
    assert {"S711-03", "S711-05"} <= set(ids)
    assert all(c["status"] == "offen" for c in ids.values())


def test_xlsx_has_header_block_numbers_and_no_formulas() -> None:
    header = {
        "report": "journal",
        "legal_entity_name": "WEG Test",
        "ledger_name": "BK",
        "period_start": date(2026, 1, 1),
        "period_end": date(2026, 12, 31),
        "as_of": None,
        "generated_at": datetime(2026, 9, 30, 10, 0, tzinfo=UTC),
        "filters": {"unit": "01"},
        "status": "draft",
        "status_note": "Entwurf",
    }
    data = report_xlsx.build_xlsx(
        "Journal",
        header,
        ["Text", "Betrag"],
        [['=HYPERLINK("http://x")', Decimal("12.50")], ["+cmd", Decimal("0.00")]],
    )
    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet is not None
    rows = [[c.value for c in r] for r in sheet.iter_rows()]
    assert rows[1][:2] == ["Rechtsträger", "WEG Test"]
    assert ["Filter", "unit=01"] in [r[:2] for r in rows]
    text_row = next(r for r in rows if r[1] == 12.5)
    assert text_row[0].startswith("'=")  # neutralised, stays a string
    cell = next(c for r in sheet.iter_rows() for c in r if c.value == 12.5)
    assert cell.number_format == "#,##0.00"
    assert all(c.data_type != "f" for r in sheet.iter_rows() for c in r)


def test_procedure_doc_marks_gaps_and_draft() -> None:
    ledger = SimpleNamespace(
        name="Buchungskreis Test",
        fiscal_year_start_month=1,
        vat_mode=VatMode.NONE,
        locked_until=None,
        leading_system=LeadingSystem.IMMOWARE24,
        migration_cutoff=None,
        template_id=None,
        template_version=None,
    )
    facts = {
        "entity": SimpleNamespace(name="WEG Test", kind=SimpleNamespace(value="hoa")),
        "posted": (3, date(2026, 1, 1), date(2026, 9, 5), 1, 1),
        "drafts": 2,
        "reversals": 0,
        "years": [(2026, 3, 3)],
        "categories": [(AccountCategory.BANK, 2)],
        "exports": [],
        "mappings": 0,
        "gates": {"G1": False, "G2": False, "G3": False, "G4": False, "G5": False},
    }
    text = procedure_doc.render(ledger, facts, datetime(2026, 9, 30, tzinfo=UTC))  # type: ignore[arg-type]
    assert "Entwurf" in text
    assert "Gebuchte Sätze: 3" in text
    assert "mit verknüpftem Originalbeleg: 1 von 3" in text
    assert "G1 Produktive Buchführung: gesperrt" in text
    assert "[zu ergänzen]" in text
    assert "\u2013" not in text
    assert "\u2014" not in text
