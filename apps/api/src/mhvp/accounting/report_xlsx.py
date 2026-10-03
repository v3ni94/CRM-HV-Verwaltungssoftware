"""Excel output (XLSX) of the journal and the evaluations (7.7, M18-05).

One workbook per evaluation: a header block with the common header fields (legal entity,
period, key date, data state, filters, draft status, M18-04), an empty row, the column titles
and the data rows. Amounts are written as numbers with the format ``#,##0.00`` (display
follows the program locale), dates as dates. Text cells pass ``csv_safe_cell`` so that no text
from the bookkeeping is interpreted as a formula. Like the CSV journal this is a neutral
export, not a DATEV or GoBD data carrier.
"""

from __future__ import annotations

import io
import uuid
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from mhvp.accounting import report_views, reports
from mhvp.accounting import services as acc
from mhvp.accounting.models import EntryStatus, JournalEntry, JournalLine, Ledger, LedgerAccount
from mhvp.core.escaping import csv_safe_cell
from mhvp.core.problems import ErrorCodes, ProblemError

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return value
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, bool):
        return "ja" if value else "nein"
    if value is None or isinstance(value, int | date):
        return value
    return csv_safe_cell(str(value))


def header_lines(header: dict[str, Any]) -> list[tuple[str, Any]]:
    """Label/value pairs of the common header in fixed order."""
    filters = header.get("filters") or {}
    return [
        ("Auswertung", header.get("report")),
        ("Rechtsträger", header.get("legal_entity_name")),
        ("Buchungskreis", header.get("ledger_name")),
        ("Zeitraum von", header.get("period_start")),
        ("Zeitraum bis", header.get("period_end")),
        ("Stichtag", header.get("as_of")),
        ("Datenstand", header.get("generated_at")),
        ("Filter", "; ".join(f"{k}={v}" for k, v in filters.items()) or "keine"),
        ("Status", "Entwurf" if header.get("status") == "draft" else header.get("status")),
        ("Hinweis", header.get("status_note")),
    ]


def build_xlsx(
    sheet_title: str,
    header: dict[str, Any],
    columns: list[str],
    rows: list[list[Any]],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover
        raise RuntimeError("workbook without sheet")
    sheet.title = sheet_title[:31]
    line = 1
    for label, value in header_lines(header):
        sheet.cell(row=line, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=line, column=2, value=_cell(value))
        line += 1
    line += 1
    for index, title in enumerate(columns, start=1):
        sheet.cell(row=line, column=index, value=title).font = Font(bold=True)
    for row in rows:
        line += 1
        for index, value in enumerate(row, start=1):
            cell = sheet.cell(row=line, column=index, value=_cell(value))
            if isinstance(value, Decimal):
                cell.number_format = MONEY_FORMAT
            elif isinstance(value, date) and not isinstance(value, datetime):
                cell.number_format = DATE_FORMAT
    for index, _title in enumerate(columns, start=1):
        width = max(
            (len(str(c.value)) for c in sheet[get_column_letter(index)] if c.value is not None),
            default=8,
        )
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 60)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


async def journal_rows(
    session: AsyncSession,
    ledger: Ledger,
    start: date,
    end: date,
    property_id: uuid.UUID | None = None,
) -> list[list[Any]]:
    rows = (
        await session.execute(
            select(JournalEntry, JournalLine, LedgerAccount)
            .join(JournalLine, JournalLine.journal_entry_id == JournalEntry.id)
            .join(LedgerAccount, LedgerAccount.id == JournalLine.account_id)
            .where(
                JournalEntry.ledger_id == ledger.id,
                JournalEntry.status == EntryStatus.POSTED,
                JournalEntry.booking_date.between(start, end),
                *([JournalLine.property_id == property_id] if property_id else []),
            )
            .order_by(JournalEntry.fiscal_year, JournalEntry.number, JournalLine.line_no)
        )
    ).all()
    return [
        [
            entry.fiscal_year,
            entry.number,
            entry.booking_date,
            entry.reference or "",
            entry.kind.value,
            line.text or entry.text,
            account.number,
            account.name,
            line.debit,
            line.credit,
            str(entry.reverses_id or ""),
            str(entry.document_id or ""),
        ]
        for entry, line, account in rows
    ]


PROPERTY_FILTER_REPORTS = frozenset({"journal", "monthly_matrix", "income_expense"})

JOURNAL_COLUMNS = [
    "Jahr",
    "Nummer",
    "Buchungstag",
    "Belegnummer",
    "Art",
    "Text",
    "Konto",
    "Kontobezeichnung",
    "Soll",
    "Haben",
    "Storno von",
    "Beleg-ID",
]

REPORTS = (
    "journal",
    "monthly_matrix",
    "target_actual",
    "trial_balance",
    "open_items",
    "revenue",
    "vat_overview",
    "income_expense",
)


LIQUIDITY_KIND_LABELS = {"free": "Frei", "reserve": "Rücklage", "deposit": "Kaution"}


async def build_report_xlsx(
    session: AsyncSession,
    ledger: Ledger,
    report: str,
    *,
    start: date,
    end: date,
    as_of: date,
    property_id: uuid.UUID | None = None,
) -> tuple[bytes, int]:
    """Builds the workbook of ``report``; returns the file and the number of data rows.

    ``property_id`` filters journal, monthly matrix and income/expense by the line's object;
    other reports have no object axis and reject the filter.
    """
    if property_id is not None and report not in PROPERTY_FILTER_REPORTS:
        raise ProblemError(
            ErrorCodes.VALIDATION,
            detail="Der Objektfilter gilt nur für Journal, Monatsmatrix und Einnahmen Ausgaben.",
        )
    if report == "journal":
        header = await report_views.report_header(
            session, ledger, report=report, start=start, end=end
        )
        rows = await journal_rows(session, ledger, start, end, property_id)
        return build_xlsx("Journal", header, JOURNAL_COLUMNS, rows), len(rows)
    if report == "monthly_matrix":
        data = await report_views.monthly_matrix(
            session, ledger, start, end, None, False, property_id
        )
        months = data["months"]
        rows = [
            [a["number"], a["name"], a["category"], *[a["months"][m] for m in months], a["total"]]
            for a in data["accounts"]
        ]
        return (
            build_xlsx(
                "Monatsmatrix",
                data["header"],
                ["Konto", "Bezeichnung", "Kategorie", *months, "Summe"],
                rows,
            ),
            len(rows),
        )
    if report == "target_actual":
        data = await report_views.target_actual(session, ledger, start, end)
        rows = [
            [
                r["number"],
                r["name"],
                r["items"],
                r["target"],
                r["actual_on_target"],
                r["difference"],
                r["receipts_in_period"],
            ]
            for r in data["rows"]
        ]
        return (
            build_xlsx(
                "Soll-Ist",
                data["header"],
                ["Konto", "Bezeichnung", "Posten", "Soll", "Ist auf Soll", "Differenz", "Eingänge"],
                rows,
            ),
            len(rows),
        )
    if report == "trial_balance":
        tb = await acc.trial_balance(session, ledger, as_of, None)
        header = await report_views.report_header(session, ledger, report=report, as_of=as_of)
        rows = [
            [a["number"], a["name"], a["category"], a["debit"], a["credit"], a["balance"]]
            for a in tb["accounts"]
        ]
        return (
            build_xlsx(
                "Saldenliste",
                header,
                ["Konto", "Bezeichnung", "Kategorie", "Soll", "Haben", "Saldo"],
                rows,
            ),
            len(rows),
        )
    if report == "open_items":
        items = await acc.open_items(session, ledger, as_of)
        header = await report_views.report_header(session, ledger, report=report, as_of=as_of)
        rows = [
            [
                i.get("account_number") or i.get("number") or "",
                i.get("kind"),
                i.get("booking_date"),
                i.get("due_date"),
                i.get("amount"),
                i.get("remaining"),
            ]
            for i in items
        ]
        return (
            build_xlsx(
                "Offene Posten",
                header,
                ["Konto", "Art", "Buchungstag", "Fällig", "Betrag", "Rest"],
                rows,
            ),
            len(rows),
        )
    if report == "revenue":
        tb = await acc.trial_balance(session, ledger, end, start)
        header = await report_views.report_header(
            session, ledger, report=report, start=start, end=end
        )
        rows = [
            [a["number"], a["name"], -a["balance"]]
            for a in tb["accounts"]
            if a["category"] == "revenue"
        ]
        return build_xlsx("Erträge", header, ["Konto", "Bezeichnung", "Betrag"], rows), len(rows)
    if report == "vat_overview":
        data = await report_views.vat_overview(session, ledger, start, end)
        rows = [
            [
                m,
                v["output_vat"],
                v["input_vat_before_deduction"],
            ]
            for m, v in data["by_month"].items()
        ]
        return (
            build_xlsx(
                "USt Entwurf",
                data["header"],
                ["Monat", "Umsatzsteuer", "Vorsteuer vor Abzugsregel"],
                rows,
            ),
            len(rows),
        )
    if report == "income_expense":
        data = await report_views.income_expense(session, ledger, start, end, False, property_id)
        rows = [["Einnahme", a["number"], a["name"], a["total"]] for a in data["revenue"]] + [
            ["Ausgabe", a["number"], a["name"], a["total"]] for a in data["cost"]
        ]
        return (
            build_xlsx(
                "Einnahmen Ausgaben",
                data["header"],
                ["Art", "Konto", "Bezeichnung", "Betrag"],
                rows,
            ),
            len(rows),
        )
    if report == "liquidity":
        data = await reports.liquidity(session, ledger, as_of)
        header = await report_views.report_header(
            session, ledger, report=report, as_of=as_of, filters={"horizon": data["horizon"]}
        )
        liq_rows: list[list[Any]] = [
            [a["number"], a["name"], LIQUIDITY_KIND_LABELS.get(a["kind"], a["kind"]), a["balance"]]
            for a in data["accounts"]
        ]
        liq_rows += [
            ["", label, "Summe", data[key]]
            for key, label in (
                ("free_funds", "Freie Mittel"),
                ("reserve_funds", "Rücklagen"),
                ("segregated_deposits", "Getrennt angelegte Kautionen"),
                ("expected_inflows", "Erwartete Einzahlungen (nicht projiziert)"),
                ("expected_outflows", "Erwartete Auszahlungen"),
                ("debtor_credits", "Guthaben von Debitoren (gebunden)"),
                ("projected_free_funds", "Freie Mittel nach Auszahlungen"),
            )
        ]
        return (
            build_xlsx("Liquidität", header, ["Konto", "Bezeichnung", "Art", "Betrag"], liq_rows),
            len(data["accounts"]),
        )
    raise ValueError(report)
